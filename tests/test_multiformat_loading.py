"""PDF / DOCX 加载与切块测试（自生成夹具，不依赖外部语料）。

夹具在 tmp_path 里现场生成，理由：
- 真实 PDF/DOCX 是二进制大文件，进仓库会持续膨胀
- 抓取外部资料容易把与项目无关的内容带进测试，且结果不可复现
- 加载器真正要防的是"文字层提不出""样式名中英文不一致"这类结构问题，
  少量可控的构造语料足以覆盖

覆盖点：
1. PDF 文字层提取 → 切块
2. DOCX 按 Heading 样式切小节（h1/h2 元数据不退化）
3. DOCX 无标题样式时退化为纯文本流程
4. 父子结构在真实文件路径上成立（parent_key 与父块正文一致）
"""

from pathlib import Path

import pytest

from app.rag.splitter import load_and_split_all, split_file


def _make_pdf(path: Path, lines: list[str]) -> None:
    """生成最小可用 PDF（单页，Helvetica，带文字层）。

    不引入 reportlab 这类重量级依赖：手写 PDF 对象结构足够让 pypdf
    提取出文字，测试关心的是"能不能提出来"，不是排版。
    """
    content = "BT /F1 12 Tf 50 750 Td 14 TL "
    content += " ".join(f"({line}) Tj T*" for line in lines)
    content += " ET"
    stream = content.encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    trailer = f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
    trailer += f"startxref\n{xref}\n%%EOF\n"
    out += trailer.encode()
    path.write_bytes(bytes(out))

def _make_docx(path: Path, with_heading: bool) -> None:
    """生成带/不带 Heading 样式的 Word 文档。

    段落内容故意超过子块上限，让一个"小节"能切出多个子块，
    才能验证父子结构（同小节子块共享 parent_key）。
    """
    from docx import Document

    doc = Document()
    body = "退货政策说明：自签收之日起七天内可申请无理由退货。" * 30
    if with_heading:
        doc.add_heading("退货政策", level=1)
        doc.add_heading("退货条件", level=2)
        doc.add_paragraph(body)
        doc.add_heading("退款时效", level=2)
        doc.add_paragraph("退款将在验收后三至五个工作日原路返回。" * 20)
    else:
        doc.add_paragraph(body)
        doc.add_paragraph("退款将在验收后三至五个工作日原路返回。" * 20)
    doc.save(str(path))

@pytest.fixture
def kb_dir(tmp_path: Path) -> Path:
    """一个含 pdf + docx 的知识库目录。"""
    _make_pdf(tmp_path / "shipping.pdf", ["Shipping takes 48 hours.", "Remote areas need 1-2 days more."])
    _make_docx(tmp_path / "returns.docx", with_heading=True)
    return tmp_path

def test_pdf_text_extracted(kb_dir: Path):
    """PDF 文字层被提取出来并切块（不是空结果）。"""
    chunks = [c for c in load_and_split_all(str(kb_dir)) if c.source.endswith(".pdf")]
    assert chunks, "PDF 加载链路未生效"
    joined = "".join(c.content for c in chunks)
    assert "Shipping takes 48 hours" in joined

def test_docx_heading_becomes_metadata(kb_dir: Path):
    """DOCX 的 Heading 样式转成 h1/h2，而不是退化为文件名。"""
    chunks = [c for c in load_and_split_all(str(kb_dir)) if c.source.endswith(".docx")]
    assert chunks
    assert {c.h1 for c in chunks} == {"退货政策"}
    assert {c.h2 for c in chunks} == {"退货条件", "退款时效"}

def test_docx_without_heading_falls_back(tmp_path: Path):
    """没有 Heading 样式的 docx 退化为纯文本流程（h1 = 文件名）。"""
    _make_docx(tmp_path / "plain.docx", with_heading=False)
    chunks = load_and_split_all(str(tmp_path))
    assert chunks
    assert all(c.h1 == "plain.docx" for c in chunks)
    assert all(c.h2 == "" for c in chunks)

def test_parent_child_structure_on_files(kb_dir: Path):
    """真实文件路径上父子结构成立：父块覆盖子块，同 parent_key 父块正文唯一。"""
    chunks = load_and_split_all(str(kb_dir))
    for c in chunks:
        assert c.parent_content
        assert len(c.parent_content) >= len(c.content)

    # 同一 (文件, parent_key) 的多个子块，父块正文必须完全一致——
    # 这是落库按 parent_key 只存一份父块的正确性前提
    by_key: dict[tuple, set] = {}
    for c in chunks:
        by_key.setdefault((c.source, c.parent_key), set()).add(c.parent_content)
    for key, parents in by_key.items():
        assert len(parents) == 1, f"{key} 的子块父块正文不一致，落库会合并错父块"

def test_split_file_matches_directory_scan(kb_dir: Path):
    """逐文件切分与目录批量切分一致（不能存在两套口径）。"""
    from app.rag.splitter import _default_strategy

    batch = load_and_split_all(str(kb_dir))
    per_file = []
    for f in sorted(kb_dir.iterdir()):
        if f.is_file():
            per_file.extend(split_file(f, _default_strategy(f.suffix.lower())))
    assert [(c.source, c.content, c.parent_key) for c in batch] == [
        (c.source, c.content, c.parent_key) for c in per_file
    ]
